# Scénář k namluvení hlasu Fyzioscanneru: všechny věty, které aplikace vyslovuje, s názvy souborů.
# Spuštění: python3 hlas/scenar.py  -> hlas/Fyzioscanner_scenar_nahravani.xlsx
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

MANUAL = "ze slovníku manuálu"
NEW = "nová formulace, prosím ověřit"

# (skupina, složka/soubor, věta, kde zazní, poznámka)
rows = []
def add(group, file, text, where, note=""):
    rows.append((group, file, text, where, note))

# --- Obecné vedení ---
add("Obecné", "obecne/pripravte-se", "Připravte se.", "Po stisku Začít cvičit, když cvik nemá vlastní úvodní pokyny.")
add("Obecné", "obecne/zacnete", "Začněte.", "Po odpočtu, začátek cvičení.")
add("Obecné", "obecne/dalsi-serie", "Další série. Začněte.", "Po pauze mezi sériemi.")
add("Obecné", "obecne/serie-hotova", "Série je hotová. Během pauzy dýchejte plynule.", "Konec série, začátek pauzy.", NEW)
add("Obecné", "obecne/hotovo", "Hotovo. Výborně, děkuji.", "Po poslední sérii.", NEW)
add("Obecné", "obecne/dokoncete-pohyb", "Dokončete pohyb dle svých možností.", "Opakování nedošlo do plného rozsahu.", NEW)
add("Obecné", "obecne/zpomalte", "Zpomalte, pohyb je plynulý.", "Pohyb byl výrazně rychlejší než ve vzoru.", MANUAL)
add("Obecné", "obecne/rozsah-vetsi", "Rozsah pohybu je větší než minule.", "Souhrn po cvičení, když se rozsah zlepšil.", NEW)
add("Obecné", "obecne/dychejte", "Dýchejte plynule do celého trupu.", "Připomínka dechu během cvičení (kamera dech nevidí).", MANUAL)

# --- Potvrzení opravy ---
for f, t in [("ano-takhle", "Ano, takhle."), ("takhle-spravne", "Takhle je to správně."), ("vyborne-udrzujte", "Výborně, tuto pozici udržujte.")]:
    add("Potvrzení", "potvrzeni/" + f, t, "Když člověk chybu opraví (střídají se).", "rozhodnuto v manuálu")

# --- Korekce (hlavní formulace) ---
K = [
 ("napremte-pater", "Napřimte páteř.", "Bederní páteř se vyhrbuje.", MANUAL),
 ("kostrc-do-dalky", "Kostrč vytáhněte do dálky.", "Bederní páteř se prohýbá.", MANUAL),
 ("hlava-v-prodlouzeni", "Hlava v prodloužení páteře.", "Záklon hlavy.", MANUAL),
 ("hlavu-ke-stropu", "Hlavu vytáhněte ke stropu.", "Předsun hlavy.", MANUAL),
 ("ramena-od-usi", "Ramena ze široka od uší.", "Ramena se zvedají k uším.", MANUAL),
 ("kolena-ve-smeru-spicek", "Kolena ve směru špiček.", "Kolena padají dovnitř.", MANUAL),
 ("kolena-nad-kotniky", "Kolena nad kotníky.", "Kolena jdou dopředu přes špičky.", NEW),
 ("ohnete-se-v-kycli", "Ohněte se v kyčelních kloubech.", "Pohyb jde z beder místo z kyčlí / málo ohnuté kyčle.", NEW),
 ("napremte-se-v-kyclich", "Napřimte se v kyčlích.", "Kyčle jsou ohnuté víc než ve vzoru.", NEW),
 ("napremte-trup", "Napřimte trup.", "Trup je nakloněný víc než ve vzoru.", NEW),
 ("trup-vpred", "Trup nakloňte vpřed.", "Trup je nakloněný méně než ve vzoru.", NEW),
 ("napremte-hrudni-pater", "Napřimte hrudní páteř.", "Hrudní páteř se vyhrbuje.", NEW),
 ("pokrcte-kolena", "Pokrčte kolena.", "Kolena jsou propnutější než ve vzoru.", NEW),
 ("kolena-mene", "Kolena pokrčte méně.", "Kolena jsou pokrčenější než ve vzoru.", NEW),
 ("pazе-predpazeni", "Paže zvedněte do předpažení.", "Paže jsou níž než ve vzoru.", NEW),
 ("paze-nize", "Paže držte níže.", "Paže jsou výš než ve vzoru.", NEW),
 ("pokrcte-lokty", "Pokrčte lokty.", "Lokty jsou propnutější než ve vzoru.", NEW),
 ("propnete-lokty", "Propněte lokty.", "Lokty jsou pokrčenější než ve vzoru.", NEW),
 ("paty-na-podlozce", "Paty zůstávají na podložce.", "Paty se odlepují.", NEW),
 ("spicky-na-podlozce", "Špičky zůstávají na podložce.", "Špičky se odlepují.", NEW),
]
for f, t, w, n in K:
    add("Korekce", "korekce/" + f.replace("е", "e"), t, w, n)

# --- Varianty (zazní místo hlavní formulace, když chyba trvá dál) ---
V = [
 ("udrzujte-siroka-ramena", "Udržujte široká ramena.", "Varianta k: Ramena ze široka od uší."),
 ("ramena-oteviena", "Ramena jsou otevřená doširoka.", "Varianta k: Ramena ze široka od uší."),
 ("hlavu-drzte-v-prodlouzeni", "Hlavu držte v prodloužení páteře.", "Varianta k: Hlava v prodloužení páteře."),
 ("hlava-krcni-pater", "Hlava je v prodloužení krční páteře.", "Varianta k: Hlava v prodloužení páteře."),
 ("hlava-zustava", "Hlava zůstává v prodloužení páteře.", "Varianta k: Hlavu vytáhněte ke stropu."),
 ("udrzujte-naprimenou", "Udržujte napřímenou páteř.", "Varianta k: Napřimte páteř. / Kostrč vytáhněte do dálky."),
 ("pater-v-cele-delce", "Páteř je napřímená v celé délce.", "Varianta k: Napřimte páteř."),
 ("kolena-nevtaci", "Kolena se nevtáčí dovnitř.", "Varianta k: Kolena ve směru špiček."),
 ("kolena-jsou-ve-smeru", "Kolena jsou ve směru špiček.", "Varianta k: Kolena ve směru špiček."),
 ("kolena-zustavaji-nad-kotniky", "Kolena zůstávají nad kotníky.", "Varianta k: Kolena nad kotníky."),
 ("pohyb-z-kycli", "Pohyb vychází z kyčelních kloubů.", "Varianta k: Ohněte se v kyčelních kloubech."),
 ("pohyb-plynuly", "Pohyb je plynulý a koordinovaný.", "Varianta k: Zpomalte, pohyb je plynulý."),
]
for f, t, w in V:
    add("Varianty", "varianty/" + f.replace("ramena-oteviena", "ramena-otevrena"), t, w, MANUAL if "kloubů" not in t and "nad kotníky" not in t else NEW)

# --- Počítání opakování ---
for i in range(1, 31):
    add("Počítání", f"pocitani/{i:02d}", str(i), "Po každém správném opakování.", "číslo vyslovte, jak běžně počítáte opakování" if i == 1 else "")

# --- Úvodní pokyny jednotlivých cviků ---
podrep = ["Postavte se, chodidla na šíři boků.", "Opora chodidel je na malíčku, kloubu pod palcem, palci a patě.", "Paže jsou volně podél těla.",
          "Hlavu vytáhněte ke stropu, ramena ze široka od uší.", "Prodechněte se do celého těla."]
add("Cvik: Dynamický podřep", "cviky/podrep/00-telefon-sikmo", "Postavte telefon šikmo před sebe, asi pod úhlem 45 stupňů.", "Úvod, varianta šikmo zepředu.", NEW)
for i, t in enumerate(podrep, 1):
    add("Cvik: Dynamický podřep", f"cviky/podrep/{i:02d}", t, "Úvodní pokyny před cvičením.", "z popisu cviku")
drep = ["Postavte se, chodidla na šíři boků.", "Opora je pod palcem, malíčkem i patou.", "Kolena lehce pokrčte.", "Udržujte kompaktně zapojenou břišní stěnu."]
for i, t in enumerate(drep, 1):
    add("Cvik: Dřep", f"cviky/drep/{i:02d}", t, "Úvodní pokyny před cvičením.", "z komentáře terapeuta, převedeno do vykání")

# --- sešit ---
wb = Workbook()
F = "Arial"
thin = Side(style="thin", color="D0D7D3")
head_fill = PatternFill("solid", fgColor="1E8C5C")
input_fill = PatternFill("solid", fgColor="FFF6CC")

# List 1: Návod
nav = wb.active; nav.title = "Návod"
lines = [
 ("Fyzioscanner – scénář k namluvení hlasu", "title"),
 ("Aplikace teď mluví hlasem telefonu. Namluvené věty ho nahradí: aplikace přehraje vaši nahrávku, a kde nahrávka chybí, použije zatím hlas telefonu. Nahrávat jde proto postupně.", ""),
 ("", ""),
 ("Jak nahrávat", "h"),
 ("• Každá věta = jeden soubor. Název souboru přesně podle sloupce „Soubor“ (složky stačí jako předpona, např. korekce_napremte-pater.wav).", ""),
 ("• Formát WAV nebo MP3, mono, 44,1 nebo 48 kHz. Tiché místo, stejná vzdálenost od mikrofonu u všech vět.", ""),
 ("• Na začátku a konci nechte asi půl vteřiny ticha, nic nestříhejte, to udělám já.", ""),
 ("• Tón jako ve voiceoverech Kinisi: klidně, věcně, vlídně. Korekce nejsou příkazy ani výtky, spíš klidné připomenutí.", ""),
 ("• Korekce jsou krátké, aby zazněly během pohybu (do 2 vteřin).", ""),
 ("• U vět označených „nová formulace, prosím ověřit“ klidně napište lepší znění do sloupce „Úprava textu“ a namluvte rovnou svou verzi.", ""),
 ("• Hotové věty označte ve sloupci „Namluveno“ (ano). Na listu Přehled se počty sečtou samy.", ""),
 ("• Hotové soubory nahrajte do sdílené složky (Google Drive) a pošlete odkaz.", ""),
 ("", ""),
 ("Co vyplnit", "h"),
 ("Žlutě podbarvené sloupce na listu Věty: „Namluveno“ a „Úprava textu“. Ostatní sloupce prosím neměňte, názvy souborů aplikace hledá přesně podle nich.", ""),
 ("", ""),
 ("Příklad vyplněného řádku", "h"),
 ("Soubor: korekce/napremte-pater  ·  Věta: Napřimte páteř.  ·  Namluveno: ano  ·  Úprava textu: (prázdné)  →  soubor korekce_napremte-pater.wav", ""),
 ("", ""),
 ("Proč tyto věty", "h"),
 ("Korekce, varianty a potvrzení vycházejí z Jazykového manuálu Kinisi (slovník chyb a ustálených pokynů). Úvodní pokyny jsou z popisů cviků. Titulky na obrazovce se nenahrávají, aplikace je jen zobrazuje.", ""),
 ("Věty s proměnným číslem jsme přeformulovali, aby stačilo nahrát je jednou (např. „Rozsah pohybu je větší než minule.“ místo čísla stupňů).", ""),
]
for i, (t, kind) in enumerate(lines, 1):
    c = nav.cell(row=i, column=1, value=t)
    c.font = Font(name=F, size=16 if kind == "title" else 11, bold=kind in ("title", "h"), color="1E8C5C" if kind == "h" else "000000")
    c.alignment = Alignment(wrap_text=True, vertical="top")
nav.column_dimensions["A"].width = 120

# List 2: Věty
ws = wb.create_sheet("Věty")
hdr = ["Č.", "Skupina", "Soubor", "Věta k namluvení", "Kde zazní", "Poznámka", "Namluveno", "Úprava textu (když jinak)"]
widths = [5, 22, 34, 46, 46, 32, 12, 40]
for j, (h, w) in enumerate(zip(hdr, widths), 1):
    c = ws.cell(row=1, column=j, value=h)
    c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = head_fill
    c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.column_dimensions[get_column_letter(j)].width = w
for i, (g, f, t, wh, n) in enumerate(rows, 2):
    vals = [i - 1, g, f, t, wh, n, "", ""]
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=i, column=j, value=v)
        c.font = Font(name=F, size=11, bold=(j == 4))
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.border = Border(bottom=thin)
        if j in (7, 8): c.fill = input_fill
        if j == 6 and n == NEW: c.font = Font(name=F, size=11, color="B45309")
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:H{len(rows)+1}"
dv = DataValidation(type="list", formula1='"ano,ne"', allow_blank=True)
ws.add_data_validation(dv); dv.add(f"G2:G{len(rows)+1}")
last = len(rows) + 1

# List 3: Přehled (vzorce)
pr = wb.create_sheet("Přehled")
groups = list(dict.fromkeys(r[0] for r in rows))
for j, h in enumerate(["Skupina", "Vět", "Namluveno", "Zbývá"], 1):
    c = pr.cell(row=1, column=j, value=h); c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = head_fill
for i, g in enumerate(groups, 2):
    pr.cell(row=i, column=1, value=g).font = Font(name=F)
    pr.cell(row=i, column=2, value=f"=COUNTIF('Věty'!$B$2:$B${last},A{i})").font = Font(name=F)
    pr.cell(row=i, column=3, value=f"=COUNTIFS('Věty'!$B$2:$B${last},A{i},'Věty'!$G$2:$G${last},\"ano\")").font = Font(name=F)
    pr.cell(row=i, column=4, value=f"=B{i}-C{i}").font = Font(name=F)
t = len(groups) + 2
pr.cell(row=t, column=1, value="Celkem").font = Font(name=F, bold=True)
for col in "BCD":
    pr[f"{col}{t}"] = f"=SUM({col}2:{col}{t-1})"; pr[f"{col}{t}"].font = Font(name=F, bold=True)
pr.column_dimensions["A"].width = 28
for col in "BCD": pr.column_dimensions[col].width = 12

from openpyxl.workbook.properties import CalcProperties
wb.calculation = CalcProperties(fullCalcOnLoad=True)  # Excel/Tabulky spočítají přehled hned při otevření
out = "hlas/Fyzioscanner_scenar_nahravani.xlsx"
wb.save(out)
print(out, len(rows), "vět")
